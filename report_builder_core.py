import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

ASSETS_DIR = os.path.join(os.path.dirname(__file__), 'report_assets')
DOCS_DIR = os.path.dirname(__file__)

# Color Palette Constants
COLOR_PRIMARY_HEX = "1E3A8A"      # Deep Navy Blue
COLOR_SECONDARY_HEX = "2563EB"    # Royal Blue
COLOR_TEAL_HEX = "0D9488"         # Deep Teal
COLOR_DARK_HEX = "1E293B"         # Charcoal Slate
COLOR_MUTED_HEX = "64748B"        # Slate Muted
COLOR_BG_LIGHT_HEX = "F1F5F9"     # Light Slate Box
COLOR_CALLOUT_BG = "EFF6FF"       # Soft Blue Callout
COLOR_BORDER_HEX = "CBD5E1"       # Light Border

COLOR_PRIMARY = RGBColor(30, 58, 138)
COLOR_SECONDARY = RGBColor(37, 99, 235)
COLOR_TEAL = RGBColor(13, 148, 136)
COLOR_DARK = RGBColor(30, 41, 59)
COLOR_MUTED = RGBColor(100, 116, 139)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=200, right=200):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m_name, m_val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m_name}')
        node.set(qn('w:w'), str(m_val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    """
    kwargs: top, bottom, left, right
    values: dict(sz=12, val='single', color='1E3A8A')
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            b_elem = OxmlElement(f'w:{edge}')
            b_elem.set(qn('w:val'), edge_data.get('val', 'single'))
            b_elem.set(qn('w:sz'), str(edge_data.get('sz', 4)))
            b_elem.set(qn('w:space'), '0')
            b_elem.set(qn('w:color'), edge_data.get('color', 'auto'))
            tcBorders.append(b_elem)
        else:
            b_elem = OxmlElement(f'w:{edge}')
            b_elem.set(qn('w:val'), 'none')
            tcBorders.append(b_elem)
    tcPr.append(tcBorders)

def add_styled_heading1(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(18)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.color.rgb = COLOR_PRIMARY
    return h

def add_styled_heading2(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(12.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_SECONDARY
    return h

def add_styled_heading3(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(3)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = COLOR_DARK
    return h

def add_body_paragraph(doc, text, bold_prefix=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(10.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_DARK
    
    r_body = p.add_run(text)
    r_body.font.name = 'Calibri'
    r_body.font.size = Pt(10.5)
    r_body.font.color.rgb = COLOR_DARK
    return p

def add_bullet_point(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_DARK
    
    r_body = p.add_run(text)
    r_body.font.name = 'Calibri'
    r_body.font.size = Pt(10)
    r_body.font.color.rgb = COLOR_DARK
    return p

def add_callout_box(doc, title, text, accent_hex=COLOR_PRIMARY_HEX, bg_hex=COLOR_CALLOUT_BG):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=160, bottom=160, left=240, right=200)
    set_cell_border(cell, 
                    left=dict(val='single', sz=24, color=accent_hex),
                    top=dict(val='single', sz=4, color=COLOR_BORDER_HEX),
                    right=dict(val='single', sz=4, color=COLOR_BORDER_HEX),
                    bottom=dict(val='single', sz=4, color=COLOR_BORDER_HEX))
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    
    r_title = p.add_run(f"📌 {title}\n")
    r_title.font.name = 'Arial'
    r_title.font.size = Pt(10)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor.from_string(accent_hex)
    
    r_text = p.add_run(text)
    r_text.font.name = 'Calibri'
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = COLOR_DARK
    
    # Space after table
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "0F172A")  # Dark editor background
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    set_cell_border(cell, 
                    left=dict(val='single', sz=12, color="38BDF8"),
                    top=dict(val='single', sz=4, color="334155"),
                    right=dict(val='single', sz=4, color="334155"),
                    bottom=dict(val='single', sz=4, color="334155"))
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    
    r = p.add_run(code_text)
    r.font.name = 'Consolas'
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(226, 232, 240)
    
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(4)

def add_figure_image(doc, img_filename, caption_text):
    img_path = os.path.join(ASSETS_DIR, img_filename)
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(6.2))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(10)
        r_cap = p_cap.add_run(caption_text)
        r_cap.font.name = 'Calibri'
        r_cap.font.size = Pt(9)
        r_cap.font.italic = True
        r_cap.font.bold = True
        r_cap.font.color.rgb = COLOR_MUTED
    else:
        print(f"Warning: Image file not found: {img_path}")

def add_styled_table(doc, headers, data, col_widths=None, header_bg=COLOR_PRIMARY_HEX):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    # Header Row
    for col_idx, h_text in enumerate(headers):
        cell = tbl.cell(0, col_idx)
        if col_widths and col_idx < len(col_widths):
            cell.width = Inches(col_widths[col_idx])
        set_cell_background(cell, header_bg)
        set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
        set_cell_border(cell,
                        top=dict(val='single', sz=6, color=header_bg),
                        bottom=dict(val='single', sz=12, color=header_bg),
                        left=dict(val='single', sz=4, color="E2E8F0"),
                        right=dict(val='single', sz=4, color="E2E8F0"))
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h_text)
        r.font.name = 'Arial'
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for row_idx, row_data in enumerate(data):
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val_text in enumerate(row_data):
            cell = tbl.cell(row_idx + 1, col_idx)
            if col_widths and col_idx < len(col_widths):
                cell.width = Inches(col_widths[col_idx])
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            set_cell_border(cell,
                            top=dict(val='single', sz=4, color=COLOR_BORDER_HEX),
                            bottom=dict(val='single', sz=4, color=COLOR_BORDER_HEX),
                            left=dict(val='single', sz=4, color=COLOR_BORDER_HEX),
                            right=dict(val='single', sz=4, color=COLOR_BORDER_HEX))
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(val_text))
            r.font.name = 'Calibri'
            r.font.size = Pt(9)
            r.font.color.rgb = COLOR_DARK

    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(6)

def create_cover_page(doc, subject_title, subject_subtitle, course_name, student_role):
    # Top banner / spacing
    p_top = doc.add_paragraph()
    p_top.paragraph_format.space_before = Pt(20)
    p_top.paragraph_format.space_after = Pt(6)
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    r_inst = p_top.add_run("ACADEMIC PROJECT SUBMISSION REPORT\nDEPARTMENT OF COMPUTER SCIENCE & ENGINEERING")
    r_inst.font.name = 'Arial'
    r_inst.font.size = Pt(11)
    r_inst.font.bold = True
    r_inst.font.color.rgb = COLOR_MUTED

    # Project Logo / Title
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(28)
    p_title.paragraph_format.space_after = Pt(8)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    r_title = p_title.add_run("HUMSAFAR (ہمسفر)")
    r_title.font.name = 'Arial'
    r_title.font.size = Pt(28)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(18)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Modern Full-Stack Travel Planning & Journey Companion Platform")
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(14)
    r_sub.font.italic = True
    r_sub.font.color.rgb = COLOR_SECONDARY

    # Subject Accent Box
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.0)
    set_cell_background(cell, "EFF6FF")
    set_cell_margins(cell, top=200, bottom=200, left=300, right=300)
    set_cell_border(cell,
                    top=dict(val='single', sz=12, color=COLOR_PRIMARY_HEX),
                    bottom=dict(val='single', sz=12, color=COLOR_PRIMARY_HEX),
                    left=dict(val='single', sz=12, color=COLOR_PRIMARY_HEX),
                    right=dict(val='single', sz=12, color=COLOR_PRIMARY_HEX))
    
    p_box = cell.paragraphs[0]
    p_box.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_subj = p_box.add_run(f"SPECIALIZED TECHNICAL REPORT:\n{subject_title.upper()}\n")
    r_subj.font.name = 'Arial'
    r_subj.font.size = Pt(13)
    r_subj.font.bold = True
    r_subj.font.color.rgb = COLOR_PRIMARY
    
    r_subj_desc = p_box.add_run(subject_subtitle)
    r_subj_desc.font.name = 'Calibri'
    r_subj_desc.font.size = Pt(10.5)
    r_subj_desc.font.color.rgb = COLOR_DARK

    # Metadata Table
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(36)

    meta_headers = ["Project Attribute", "Specification Details"]
    meta_data = [
        ["Project Title", "HUMSAFAR: Comprehensive Trip Planner & Companion"],
        ["Target Course", course_name],
        ["Academic Focus", student_role],
        ["Submission Year", "Academic Session 2026"],
        ["System Core Stack", "React 18, Vite, Node.js, Express 5, SQLite, Leaflet, Firestore"],
        ["Project Status", "Fully Implemented, Verified & End-to-End Automated Tested (100% Pass)"]
    ]
    add_styled_table(doc, meta_headers, meta_data, col_widths=[2.2, 4.3])

    doc.add_page_break()

print("Helper definitions loaded successfully.")
